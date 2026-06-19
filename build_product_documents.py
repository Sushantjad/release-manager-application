from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "documents"
OUT.mkdir(exist_ok=True)

BLUE = "1F4E78"
DARK_BLUE = "17365D"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FA"
GRAY = "667085"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D0D5DD"
TEXT = "1F2937"
GREEN = "16803C"
AMBER = "B54708"
RED = "B42318"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table_pr.append(tbl_w)
    total = int(sum(widths) * 1440)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = table_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        table_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_document(doc, running_title):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (11.5, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.1

    header = section.header
    p = header.paragraphs[0]
    p.text = running_title
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    run = p.runs[0]
    run.font.name = "Calibri"
    run.font.size = Pt(8.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(GRAY)

    footer = section.footer
    p = footer.paragraphs[0]
    add_page_number(p)


def add_title_block(doc, title, subtitle, version="Version 1.0", status="Draft for stakeholder review"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(25)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(subtitle)
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(GRAY)

    table = doc.add_table(rows=4, cols=2)
    set_table_width(table, [1.5, 5.0])
    set_table_borders(table, color="D8E2EB", size="4")
    metadata = [
        ("Document", title),
        ("Version", version),
        ("Status", status),
        ("Prepared", date.today().strftime("%d %B %Y")),
    ]
    for idx, (label, value) in enumerate(metadata):
        left, right = table.rows[idx].cells
        set_cell_shading(left, PALE_BLUE)
        left.text = label
        right.text = value
        left.paragraphs[0].runs[0].bold = True
        left.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(DARK_BLUE)
        for cell in (left, right):
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    doc.add_paragraph()


def add_callout(doc, label, text, color=BLUE, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.5])
    set_table_borders(table, color=fill, size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor.from_string(TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_table(doc, headers, rows, widths, header_fill=BLUE, small=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        cell.text = header
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
            run.font.size = Pt(8.5 if small else 9.5)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = str(value)
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                run.font.size = Pt(8.5 if small else 9.5)
        set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_requirement(doc, req_id, title, statement, acceptance=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{req_id} | {title}")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p = doc.add_paragraph(statement)
    if acceptance:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("Acceptance:")
        r.bold = True
        for item in acceptance:
            q = doc.add_paragraph(style="List Bullet")
            q.add_run(item)


def build_prd():
    doc = Document()
    configure_document(doc, "Release Manager Application | Product Requirements")
    add_title_block(
        doc,
        "Product Requirements Document",
        "Release Manager Application - end-to-end release planning and command center",
    )
    add_callout(
        doc,
        "Product vision",
        "Provide one trusted workspace for planning, governing, communicating and closing application releases, reducing manual coordination while improving management visibility.",
    )

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "Release planning currently depends on Jira updates, spreadsheets, messages, emails and repeated follow-ups across project managers, release planners, application teams and management. This creates duplicated effort, inconsistent status information and limited visibility on release day."
    )
    doc.add_paragraph(
        "The Release Manager Application will connect release scope from Jira with RMI readiness, deployment schedules, incident tracking and management communication. Version 1 focuses on a dependable operational workflow and a clear command-center view."
    )

    doc.add_heading("2. Problem Statement", level=1)
    add_bullets(
        doc,
        [
            "Release scope and RMI readiness are manually reconciled across teams.",
            "Release planners repeatedly contact Scrum Masters, developers and testers for UAT status.",
            "Planned and actual deployment times are updated manually with limited consistency.",
            "Management status emails require hourly manual preparation and distribution.",
            "Jira incidents are difficult to consolidate into a single release view.",
            "Release closure lacks a consistent summary and audit trail.",
        ],
    )

    doc.add_heading("3. Goals and Success Measures", level=1)
    add_table(
        doc,
        ["Goal", "Success measure for V1"],
        [
            ("Single source of truth", "All in-scope RMIs, applications, deployments and incidents are visible under one release."),
            ("Reduce manual reporting", "Release status statement can be drafted and distributed from the application."),
            ("Improve readiness visibility", "Every RMI has an owner and traffic-light readiness status."),
            ("Improve release-day control", "Planned and actual deployment times plus blockers are visible on one dashboard."),
            ("Strengthen governance", "Key status, assignment and timing changes are recorded in an audit history."),
        ],
        [2.2, 4.3],
    )

    doc.add_heading("4. Users and Roles", level=1)
    add_table(
        doc,
        ["Role", "Primary responsibilities", "Access"],
        [
            ("Project Manager", "Own project requirements, monitor RMI readiness, assign release planners.", "Create/update assigned planning records; view release execution."),
            ("Release Manager", "Own release setup, governance, deployment command center, communications and closure.", "Full operational access for assigned releases."),
            ("Release Planner", "Coordinate UAT readiness and update RMI status.", "Update assigned RMIs; view related release and deployment data."),
            ("Intake Team", "Review and approve RMI tagging; associate RMIs with a release in Jira.", "Integration-driven scope; optional review access."),
            ("Business / Management", "Monitor readiness, progress, incidents and outcome.", "Read-only access."),
            ("Application Owner / Deployment POC", "Provide CR and deployment details; confirm actual start/end.", "Update assigned deployment records."),
        ],
        [1.2, 3.1, 2.2],
        small=True,
    )

    doc.add_heading("5. End-to-End Workflow", level=1)
    add_numbered(
        doc,
        [
            "Project Manager gathers stakeholder requirements and creates or updates the Jira item and RMI.",
            "Intake Team reviews the RMI, confirms approval and tags it to a specific release in Jira.",
            "Jira synchronization creates or updates the release scope in the Release Manager Application.",
            "Project Manager assigns each RMI to a Release Planner.",
            "Release Planner coordinates UAT and marks the RMI Orange, Red or Green.",
            "Application Owner submits deployment details including version, change request, planned window, POC and remarks.",
            "Release Manager validates readiness, configures distribution lists and starts the release.",
            "Teams update deployment actual times, incidents and release health during execution.",
            "Release Manager sends hourly status communication and escalates blockers.",
            "After all deployment work is completed or dispositioned, Release Manager stops and closes the release.",
        ],
    )

    doc.add_heading("6. Product Scope", level=1)
    doc.add_heading("6.1 Version 1 In Scope", level=2)
    add_bullets(
        doc,
        [
            "Dashboard with release controls, application counts, Jira incident count and management summary.",
            "Release register containing RMI, project, ownership and application details.",
            "RMI readiness management with planner assignment and Orange/Red/Green status.",
            "Deployment register with planned and actual times, CR number, POC and remarks.",
            "Release-day command center with Start/Stop, health, distribution lists and status statements.",
            "Jira synchronization contract for releases, RMIs and incidents using Affect Version / release identifier.",
            "Role-based access, validation, audit history, filters and export.",
            "AI-assisted draft summaries, risk highlights and post-release recap, requiring human approval.",
        ],
    )
    doc.add_heading("6.2 Out of Scope for Version 1", level=2)
    add_bullets(
        doc,
        [
            "Executing application deployments directly from the platform.",
            "Replacing Jira as the system of record for requirements or incident ownership.",
            "Autonomous AI decisions, automatic release cancellation or automatic blocker resolution.",
            "Advanced predictive models trained on enterprise historical data.",
            "Native mobile applications and complex portfolio financial management.",
        ],
    )

    doc.add_heading("7. Functional Capabilities", level=1)
    add_table(
        doc,
        ["Capability", "Purpose", "Key outcome"],
        [
            ("Dashboard", "Summarize release health and execution.", "Fast management and operator visibility."),
            ("Release Management", "Define release scope and ownership.", "Trusted release register."),
            ("RMI Readiness", "Track UAT readiness and blockers.", "Clear go/no-go preparation."),
            ("Deployment Planning", "Capture schedule, CR and actual timings.", "Controlled deployment sequence."),
            ("Release Day", "Start, monitor, communicate and stop a release.", "Consistent command-center operation."),
            ("Jira Integration", "Synchronize approved scope and incidents.", "Less duplicate data entry."),
            ("Communications", "Generate and send status updates to DLs.", "Reduced hourly reporting effort."),
            ("AI Assistance", "Draft summaries and identify attention areas.", "Faster, more consistent decisions."),
        ],
        [1.5, 2.6, 2.4],
        small=True,
    )

    doc.add_heading("8. Business Rules", level=1)
    add_bullets(
        doc,
        [
            "RMI Orange means UAT is in progress; Red means a blocker exists; Green means UAT completed with no open blocking issue.",
            "A release cannot be marked Ready when any in-scope RMI is Red or lacks a Release Planner.",
            "Deployment actual start and end times must not overwrite planned times.",
            "A deployment cannot be completed without an actual end time and final remark or outcome.",
            "Release health is Green for smooth progress, Yellow for progress with issues, and Red when stopped or materially blocked.",
            "Closing a release requires every deployment to be Completed, Cancelled or Deferred with a recorded disposition.",
            "AI-generated content must be visibly marked as a draft and approved by an authorized user before sending.",
        ],
    )

    doc.add_heading("9. Integration Requirements", level=1)
    doc.add_heading("9.1 Jira", level=2)
    add_bullets(
        doc,
        [
            "Import release, RMI and project data after Intake approval and release tagging.",
            "Use Jira issue key as the stable external identifier and avoid duplicate records.",
            "Import incidents when Affect Version matches the configured release name or identifier.",
            "Display synchronization status, timestamp and errors without blocking unrelated application work.",
            "Define field mappings, permissions, API authentication and synchronization frequency during technical design.",
        ],
    )
    doc.add_heading("9.2 Email / Distribution Lists", level=2)
    add_bullets(
        doc,
        [
            "Store one or more approved distribution lists for each release.",
            "Allow preview and editing of every status message before sending.",
            "Record sender, recipients, timestamp, subject and delivery result.",
            "Use the organization's approved email service and security controls.",
        ],
    )

    doc.add_heading("10. AI-Assisted Capabilities", level=1)
    add_table(
        doc,
        ["AI feature", "Inputs", "Output", "Control"],
        [
            ("Status summary draft", "RMI, deployment and incident statuses", "Hourly management statement", "Release Manager reviews and sends."),
            ("Risk highlights", "Red RMIs, delays, open incidents, remarks", "Prioritized attention list", "Advisory only; source records linked."),
            ("Post-release recap", "Timeline, deployment outcomes, incidents", "Draft closure summary", "Human approval before publication."),
            ("Data quality prompt", "Missing or inconsistent fields", "Suggested corrections", "User accepts or rejects each change."),
        ],
        [1.4, 2.0, 1.9, 1.2],
        small=True,
    )

    doc.add_heading("11. Non-Functional Requirements", level=1)
    add_table(
        doc,
        ["Area", "Requirement"],
        [
            ("Security", "Enterprise authentication, least-privilege role access and encrypted transport."),
            ("Availability", "Available throughout planned release windows with monitored service health."),
            ("Performance", "Primary views should load within 3 seconds for normal release volumes."),
            ("Auditability", "Material changes record user, time, previous value and new value."),
            ("Accessibility", "Keyboard navigation, meaningful labels and non-color status indicators."),
            ("Data integrity", "Validation, stable identifiers and safe synchronization retries."),
            ("Privacy", "No unnecessary personal or confidential data in AI prompts or exports."),
        ],
        [1.3, 5.2],
    )

    doc.add_heading("12. Assumptions and Dependencies", level=1)
    add_bullets(
        doc,
        [
            "Jira provides approved API access and consistent field mappings for RMI, release and incident records.",
            "The organization confirms authoritative ownership for project, release, planner and application data.",
            "An approved email integration and distribution-list policy are available.",
            "Release status terminology and closure controls are agreed by Release Management and business stakeholders.",
            "AI services, if enabled, meet enterprise security, privacy and retention requirements.",
        ],
    )

    doc.add_heading("13. Risks and Mitigations", level=1)
    add_table(
        doc,
        ["Risk", "Impact", "Mitigation"],
        [
            ("Poor Jira data quality", "Incomplete or misleading release scope", "Field validation, sync errors and ownership dashboard."),
            ("Low adoption", "Teams continue parallel spreadsheets", "Simple workflow, role training and agreed source-of-truth policy."),
            ("Late status updates", "Management view becomes stale", "Timestamp visibility, reminders and overdue indicators."),
            ("AI summary error", "Incorrect communication", "Source-linked drafts and mandatory human approval."),
            ("Integration outage", "Delayed synchronization", "Retry queue, visible last-sync status and controlled manual fallback."),
        ],
        [2.0, 2.1, 2.4],
        small=True,
    )

    doc.add_heading("14. Release and Adoption Plan", level=1)
    add_numbered(
        doc,
        [
            "Validate workflow and terminology with Project Managers, Release Managers and business readers.",
            "Pilot V1 with one release and a small set of applications.",
            "Measure completeness, update timeliness, reporting effort and user feedback.",
            "Resolve process gaps and finalize Jira/email integration controls.",
            "Expand to additional releases with role-based onboarding and governance.",
        ],
    )

    doc.add_heading("15. Open Decisions", level=1)
    add_bullets(
        doc,
        [
            "Which Jira issue types and fields represent RMI, approval, release and application?",
            "Which system owns the final application inventory and contact details?",
            "What is the required hourly communication schedule and template?",
            "Who may change overall release health to Red and who may restart a stopped release?",
            "What retention period applies to audit records and release communications?",
            "Which AI platform is approved for enterprise release data?",
        ],
    )

    path = OUT / "Release_Manager_Product_Requirements_Document_v1.0.docx"
    doc.save(path)
    return path


def build_feature_spec():
    doc = Document()
    configure_document(doc, "Release Manager Application | Version 1 Feature Specification")
    add_title_block(
        doc,
        "Version 1 Feature Specification",
        "Release Manager Application - functional behavior, data, rules and acceptance criteria",
    )
    add_callout(
        doc,
        "V1 objective",
        "Enable a Release Manager to prepare, run, communicate and close an application release from a single controlled workspace.",
    )

    doc.add_heading("1. Specification Conventions", level=1)
    add_bullets(
        doc,
        [
            "Must indicates a Version 1 requirement.",
            "Should indicates a desirable behavior that may be deferred only through product approval.",
            "Acceptance criteria use observable outcomes and are suitable for testing.",
            "All timestamps are stored consistently and displayed in the release's configured time zone.",
        ],
    )

    doc.add_heading("2. Roles and Permission Matrix", level=1)
    add_table(
        doc,
        ["Action", "Release Manager", "Project Manager", "Release Planner", "App Owner / POC", "Business"],
        [
            ("View release data", "Yes", "Yes", "Assigned", "Assigned", "Read"),
            ("Create/edit release", "Yes", "Limited", "No", "No", "No"),
            ("Assign RMI planner", "Yes", "Yes", "No", "No", "No"),
            ("Update RMI status", "Yes", "Yes", "Assigned", "No", "No"),
            ("Update deployment", "Yes", "View", "View", "Assigned", "No"),
            ("Start/stop release", "Yes", "No", "No", "No", "No"),
            ("Send status email", "Yes", "No", "No", "No", "No"),
            ("View audit/export", "Yes", "Yes", "Assigned", "Assigned", "Read"),
        ],
        [1.55, 1.05, 1.05, 1.05, 1.15, 0.65],
        small=True,
    )

    doc.add_heading("3. Release Lifecycle", level=1)
    add_table(
        doc,
        ["State", "Meaning", "Permitted transition"],
        [
            ("Draft", "Release created; scope and schedule incomplete.", "Planning"),
            ("Planning", "RMIs assigned and deployment details collected.", "Ready or Draft"),
            ("Ready", "Required readiness checks passed.", "In Progress or Planning"),
            ("In Progress", "Release Start has been recorded.", "Stopped or Completed"),
            ("Stopped", "Release execution halted because of a blocker.", "In Progress or Completed"),
            ("Completed", "All deployments dispositioned and release closed.", "No normal transition"),
        ],
        [1.15, 3.2, 2.15],
    )

    doc.add_heading("4. Dashboard", level=1)
    add_requirement(
        doc,
        "DASH-001",
        "Release summary",
        "The dashboard must display the selected release name, date, environment, lifecycle state, overall health and last updated time.",
        [
            "Selecting a release refreshes all dashboard values.",
            "Business users can view but cannot change dashboard values.",
        ],
    )
    add_requirement(
        doc,
        "DASH-002",
        "Start and Stop controls",
        "The dashboard must provide Start Release and Stop Release controls to authorized Release Managers.",
        [
            "Start records the actual release start time and changes the lifecycle to In Progress.",
            "Stop requires a reason, records the time and changes lifecycle to Stopped.",
            "Controls are disabled when the lifecycle does not allow the action.",
        ],
    )
    add_requirement(
        doc,
        "DASH-003",
        "Operational metrics",
        "The dashboard must show Completed, Ongoing and Pending application counts plus the Jira incident count.",
        [
            "Counts match the underlying deployment and incident lists.",
            "Selecting a metric filters or opens the corresponding records.",
        ],
    )
    add_requirement(
        doc,
        "DASH-004",
        "Management status",
        "The dashboard must display the latest approved management status statement and its sent time.",
        [
            "Draft statements are not shown as sent updates.",
            "The latest approved statement is readable without edit access.",
        ],
    )

    doc.add_heading("5. Release Page", level=1)
    add_table(
        doc,
        ["Column", "Required", "Behavior"],
        [
            ("RMI Number", "Yes", "Unique within the release; links to RMI detail."),
            ("Project Details", "Yes", "Project or feature summary synchronized from Jira where available."),
            ("Project Manager", "Yes", "Named accountable PM."),
            ("Release Manager", "Yes", "Named owner for the release."),
            ("Application Name", "Yes", "One row per RMI/application association."),
        ],
        [1.5, 1.0, 4.0],
    )
    add_requirement(
        doc,
        "REL-001",
        "Jira scope synchronization",
        "Approved RMIs tagged to the release in Jira must appear on the Release page without manual re-entry.",
        [
            "A repeated sync updates the same record rather than creating a duplicate.",
            "Removed or changed Jira associations are visibly reconciled according to the agreed retention rule.",
            "Last synchronization time and errors are available.",
        ],
    )
    add_requirement(
        doc,
        "REL-002",
        "Search, filter and export",
        "Users must be able to search by RMI, project, person and application; filter by owner/status; and export the current result.",
        [
            "Export respects the active filters.",
            "Read-only users can export only fields they are authorized to view.",
        ],
    )

    doc.add_heading("6. RMI Readiness Page", level=1)
    add_table(
        doc,
        ["Column", "Required", "Behavior"],
        [
            ("RMI Number", "Yes", "Stable key linked to release and Jira issue."),
            ("RMI Status", "Yes", "Orange, Red or Green plus a text label."),
            ("Release Planner", "Yes", "Assigned responsible planner."),
        ],
        [1.5, 1.0, 4.0],
    )
    add_requirement(
        doc,
        "RMI-001",
        "Planner assignment",
        "A Project Manager or Release Manager must be able to assign a Release Planner to each RMI.",
        [
            "Assignment records the acting user and timestamp.",
            "Unassigned RMIs are highlighted and prevent Ready status.",
        ],
    )
    add_requirement(
        doc,
        "RMI-002",
        "Readiness status",
        "An authorized user must update readiness as Orange - UAT In Progress, Red - Blocked, or Green - UAT Complete.",
        [
            "Red requires blocker details and an owner.",
            "Green requires UAT completion confirmation.",
            "Status is represented by color, text and an accessible icon or label.",
        ],
    )
    add_requirement(
        doc,
        "RMI-003",
        "Readiness history",
        "The application must retain a history of RMI status, planner and blocker changes.",
        [
            "History contains previous value, new value, user, timestamp and comment when supplied.",
        ],
    )

    doc.add_heading("7. Deployment Page", level=1)
    add_table(
        doc,
        ["Field / Column", "Required", "Validation or behavior"],
        [
            ("RMI Number", "Yes", "Must reference an RMI in the selected release."),
            ("Application Name", "Yes", "Must reference the RMI/application scope."),
            ("Version", "Yes", "Text value appropriate to the application."),
            ("Planned Start Time", "Yes", "Must precede planned end time."),
            ("Planned End Time", "Yes", "Must follow planned start time."),
            ("Actual Start Time", "Conditional", "Required after deployment is started."),
            ("Actual End Time", "Conditional", "Required when deployment is completed."),
            ("CR Number", "Yes", "Change request identifier; uniqueness warning if reused."),
            ("Release-night POC", "Yes", "Named contact with approved contact reference."),
            ("Remark", "Yes", "Current operational note or final outcome."),
        ],
        [1.45, 1.05, 4.0],
        small=True,
    )
    add_requirement(
        doc,
        "DEP-001",
        "Create and maintain deployment plan",
        "Authorized users must create and edit deployment records using the required fields above.",
        [
            "The application prevents saving when mandatory data is missing or time order is invalid.",
            "Changes to planned times are audited.",
        ],
    )
    add_requirement(
        doc,
        "DEP-002",
        "Deployment execution status",
        "Each deployment must support Pending, In Progress, Completed, Blocked, Cancelled and Deferred states.",
        [
            "Starting a deployment defaults actual start to the current time but allows correction with audit.",
            "Completing a deployment requires actual end and a final remark.",
            "Blocked requires blocker details and updates dashboard health indicators.",
        ],
    )
    add_requirement(
        doc,
        "DEP-003",
        "Schedule variance",
        "The application should show late start, overrunning and duration variance indicators.",
        [
            "A pending deployment after planned start is flagged as late.",
            "An in-progress deployment after planned end is flagged as overrunning.",
        ],
    )

    doc.add_heading("8. Release Day Command Center", level=1)
    add_requirement(
        doc,
        "CMD-001",
        "Distribution lists",
        "The Release Manager must add one or more distribution lists before the first status email is sent.",
        [
            "Invalid email formats are rejected.",
            "Recipients are visible in the preview before sending.",
        ],
    )
    add_requirement(
        doc,
        "CMD-002",
        "Overall release health",
        "The Release Manager must set Green - Smooth, Yellow - Proceeding with Issues, or Red - Stopped / Blocked.",
        [
            "Yellow and Red require a status explanation.",
            "Every health change is timestamped and audited.",
        ],
    )
    add_requirement(
        doc,
        "CMD-003",
        "Hourly status statement",
        "The application must allow the Release Manager to create, preview, edit and send a management status statement.",
        [
            "The statement contains release health, deployment progress, incidents, blockers and next checkpoint.",
            "The system records sender, recipients, subject, body, sent time and delivery result.",
            "A reminder is displayed when the configured update interval is overdue.",
        ],
    )
    add_requirement(
        doc,
        "CMD-004",
        "Release closure",
        "The Release Manager must be able to close the release after all deployment records have a final disposition.",
        [
            "Closure records actual end time and final summary.",
            "Unresolved records are listed and block closure unless an authorized override with reason is used.",
            "Completed release data becomes read-only except for authorized administrative correction.",
        ],
    )

    doc.add_heading("9. Jira Incident Integration", level=1)
    add_requirement(
        doc,
        "JIRA-001",
        "Incident import",
        "Jira incident tickets whose Affect Version matches the release identifier must appear in the Release Manager Application.",
        [
            "Each incident displays key, summary, priority, status, assignee and last updated time.",
            "Repeated synchronization updates existing incidents without duplication.",
            "The dashboard Jira count matches active incident filter rules.",
        ],
    )
    add_requirement(
        doc,
        "JIRA-002",
        "Synchronization resilience",
        "A Jira synchronization failure must be visible and must not erase previously synchronized data.",
        [
            "Users see the last successful synchronization time.",
            "Authorized users can retry synchronization.",
            "Errors are logged with enough detail for support without exposing secrets.",
        ],
    )

    doc.add_heading("10. AI Assistance", level=1)
    add_requirement(
        doc,
        "AI-001",
        "Management summary draft",
        "The system should generate an editable status draft from current release data.",
        [
            "The draft distinguishes facts from suggested narrative.",
            "It links or references the source records used.",
            "It cannot be sent without Release Manager review and confirmation.",
        ],
    )
    add_requirement(
        doc,
        "AI-002",
        "Risk attention list",
        "The system should highlight blocked RMIs, delayed deployments, high-priority incidents and stale updates.",
        [
            "Every highlight states the underlying reason.",
            "Users can dismiss or acknowledge a highlight without altering source data.",
        ],
    )
    add_requirement(
        doc,
        "AI-003",
        "Post-release summary",
        "The system should draft a closure recap containing planned versus actual timing, deployment outcomes, incidents and follow-up actions.",
        [
            "The draft is editable and not published automatically.",
            "Sensitive data is handled according to enterprise AI policy.",
        ],
    )

    doc.add_heading("11. Notifications and Reminders", level=1)
    add_table(
        doc,
        ["Trigger", "Recipient", "V1 behavior"],
        [
            ("RMI assigned", "Release Planner", "In-app notification; optional email."),
            ("RMI becomes Red", "PM and Release Manager", "Immediate alert with blocker details."),
            ("Deployment nearing planned start", "Deployment POC", "Configurable reminder."),
            ("Deployment late or overrunning", "Release Manager and POC", "Command-center warning."),
            ("Hourly status overdue", "Release Manager", "Persistent reminder until sent or deferred."),
            ("Jira sync failure", "Release Manager / support", "Visible warning and retry action."),
        ],
        [2.1, 1.8, 2.6],
        small=True,
    )

    doc.add_heading("12. Audit, Search and Reporting", level=1)
    add_bullets(
        doc,
        [
            "Audit history must cover lifecycle changes, assignments, readiness, planned/actual times, health, communications and closure.",
            "Users must filter by release, application, owner, RMI status, deployment status and incident status.",
            "V1 exports must support a readable spreadsheet or CSV representation of the selected view.",
            "A release summary must include planned versus actual timing, completion counts, incidents and final outcome.",
        ],
    )

    doc.add_heading("13. Core Data Model", level=1)
    add_table(
        doc,
        ["Entity", "Key fields"],
        [
            ("Release", "ID, name, date, environment, manager, lifecycle, health, planned/actual start/end, time zone"),
            ("RMI", "ID, Jira key, project details, PM, planner, readiness, blocker, release ID"),
            ("Application", "ID, name, service, environment, owner"),
            ("Deployment", "ID, release, RMI, application, version, CR, planned/actual times, status, POC, remark"),
            ("Incident", "Jira key, release, summary, priority, status, assignee, updated time"),
            ("Communication", "release, recipients, subject, body, author, sent time, delivery result"),
            ("Audit Event", "entity, entity ID, action, old/new value, user, timestamp, reason"),
        ],
        [1.35, 5.15],
        small=True,
    )

    doc.add_heading("14. Global Validation and Experience Rules", level=1)
    add_bullets(
        doc,
        [
            "Required fields are clearly marked and validation messages explain how to correct the value.",
            "Status is never communicated by color alone.",
            "Tables support sorting, filtering, pagination and a useful empty state.",
            "Dates and times display the release time zone and use a consistent format.",
            "Unsaved changes prompt the user before navigation.",
            "Read-only users see disabled or absent editing controls.",
            "Every screen shows loading, success, empty and error states.",
        ],
    )

    doc.add_heading("15. V1 End-to-End Acceptance Scenario", level=1)
    add_numbered(
        doc,
        [
            "An approved Jira RMI tagged to Release R1 synchronizes into the Release page.",
            "A Project Manager assigns the RMI to a Release Planner.",
            "The planner changes the RMI from Orange to Green after UAT completion.",
            "An Application Owner records version, CR, planned window, POC and remark.",
            "The Release Manager confirms readiness, distribution lists and starts R1.",
            "The deployment POC starts and completes deployment with actual times.",
            "A Jira incident with Affect Version R1 appears and updates the incident count.",
            "The Release Manager generates, edits and sends an hourly status update.",
            "All deployments receive a final disposition and the Release Manager closes R1.",
            "The system retains the release summary, communications and audit history as read-only records.",
        ],
    )

    doc.add_heading("16. Definition of Done for Version 1", level=1)
    add_bullets(
        doc,
        [
            "All Must requirements in this specification pass functional acceptance testing.",
            "Role permissions are tested for every supported role.",
            "Jira synchronization is validated against agreed test fields and failure scenarios.",
            "Email preview, sending and audit behavior are validated using an approved test service.",
            "Accessibility, performance, security and audit requirements pass agreed checks.",
            "Operational support, ownership and recovery procedures are documented.",
            "Pilot users complete one representative release workflow and approve production readiness.",
        ],
    )

    doc.add_heading("17. Traceability Summary", level=1)
    add_table(
        doc,
        ["Business need", "Feature requirements"],
        [
            ("Management visibility", "DASH-001 to DASH-004, CMD-002, CMD-003"),
            ("Less manual planning", "REL-001, RMI-001 to RMI-003, DEP-001"),
            ("Release-day control", "DASH-002, DEP-002, DEP-003, CMD-001 to CMD-004"),
            ("Jira consolidation", "REL-001, JIRA-001, JIRA-002"),
            ("Faster reporting", "CMD-003, AI-001, AI-003"),
            ("Governance and accountability", "RMI-003, CMD-004, audit requirements"),
        ],
        [2.4, 4.1],
    )

    path = OUT / "Release_Manager_V1_Feature_Specification_v1.0.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    prd = build_prd()
    spec = build_feature_spec()
    print(prd)
    print(spec)
